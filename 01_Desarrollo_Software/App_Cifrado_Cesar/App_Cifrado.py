import tkinter as tk
from tkinter import filedialog
import docx
from pypdf import PdfReader

# --- 1.1. FUNCIONES ---
# --- 1.1. LÓGICA DE CIFRADO Y DESCIFRADO ---
def cifrar_cesar(Texto, k):
    Texto_cifrado = ""
    for letra in Texto:
        if letra.isupper():
            X = ord(letra)
            Xn = X - 65
            Xcif = Xn + k
            Texto_cifrado += chr((Xcif % 26) + 65)
        elif letra.islower():
            X = ord(letra)
            Xn = X - 97
            Xcif = Xn + k
            Texto_cifrado += chr((Xcif % 26) + 97)
        else:
            Texto_cifrado += letra
    return Texto_cifrado

def descifrar_cesar(Texto, k):
    Texto_descifrado = ""
    for letra in Texto:
        if letra.isupper():
            X = ord(letra)
            Xn = X - 65
            Xcif = Xn - k
            Texto_descifrado += chr((Xcif % 26) + 65)
        elif letra.islower():
            X = ord(letra)
            Xn = X - 97
            Xcif = Xn - k
            Texto_descifrado += chr((Xcif % 26) + 97)
        else:
            Texto_descifrado += letra
    return Texto_descifrado

# --- 1.2. LIMPIEZA DE CAJAS DE TEXTO ---
def limpiar_campos():
    entrada_texto.delete("1.0", tk.END)
    entrada_clave.delete(0, tk.END)
    entrada_resultado.config(state="normal")
    entrada_resultado.delete("1.0", tk.END)
    entrada_resultado.config(state="disabled")

# --- 1.3. COPIADO RÁPIDO ---
def copiar_resultado():
    texto_a_copiar = entrada_resultado.get("1.0", tk.END).strip()
    if texto_a_copiar and not texto_a_copiar.startswith("!"):
        ventana.clipboard_clear
        ventana.clipboard_append(texto_a_copiar)

# --- 1.4. CIFRAR ARCHIVO TEXTO---
def cifrar_archivo():
    try:
        clave=int(entrada_clave.get())
    except ValueError:
        entrada_resultado.config(state="normal")
        entrada_resultado.delete("1.0", tk.END)
        entrada_resultado.insert(tk.END, "!Ingrese una clave númerica válida!")
        entrada_resultado.config(state="disabled")
        return
    ruta_archivo = filedialog.askopenfilename(
        title="Seleccona un archivo para cifrar",
        filetypes=[
            ("Archivos compatibles", "*.txt *.docx *.pdf"), 
            ("Archivos de Word", "*.docx"),
            ("Archivos PDF", "*.pdf"),
            ("Todos los archivos", "*.*")
        ]
    )
    if not ruta_archivo:
        return
    texto_completo = ""
    try:
        if ruta_archivo.endswith(".txt"):
            with open(ruta_archivo,"r",encoding="utf-8") as f:
                texto_completo = f.read()
        elif ruta_archivo.endswith(".docx"):
            doc = docx.Document(ruta_archivo)
            texto_completo = "\n".join([p.text for p in doc.paragraphs])
        elif ruta_archivo.endswith(".pdf"):
            reader = PdfReader(ruta_archivo)
            for pagina in reader.pages:
                texto_extraido = pagina.extract_text()
                if texto_extraido:
                    texto_completo += texto_extraido + "\n"

        texto_cifrado = cifrar_cesar(texto_completo,clave)
        
        nombre_sugerido = ruta_archivo.rsplit(".",1)[0] + "_cifrado.txt"
        ruta_salida = filedialog.asksaveasfilename(
            title="Guardar archivo cifrado como...",
            initialfile=nombre_sugerido.split("/")[-1],
            defaultextension=".txt",
            filetypes=[
                ("Archivo de texto", "*.txt"),
                ("Archivo de Word", "*.docx"),
                ("Todos los archivos", "*.*")
                ]
        )
        if not ruta_salida:
            return

        if ruta_salida.endswith(".docx"):
            doc_salida = docx.Document()
            doc_salida.add_paragraph(texto_cifrado)
            doc_salida.save(ruta_salida)
        else:
            with open(ruta_salida,"w",encoding="utf-8") as f:
                f.write(texto_cifrado)

        entrada_resultado.config(state="normal")
        entrada_resultado.delete("1.0", tk.END)
        entrada_resultado.insert(tk.END, f"¡Archivo cifrado con éxito!\nGuardado en:\n{ruta_salida}")
        entrada_resultado.config(state="disabled")

    except Exception as e:
        entrada_resultado.config(state="normal")
        entrada_resultado.delete("1.0", tk.END)
        entrada_resultado.insert(tk.END, f"Error al procesar el archivo: {e}")
        entrada_resultado.config(state="disabled")

# --- 2. CONEXIÓN ENTRE LA INTERFAZ Y LAS FUNCIONES ---
def ejecutar_cifrado():
    texto = entrada_texto.get("1.0", tk.END).strip()
    try:
        clave = int(entrada_clave.get())
        resultado = cifrar_cesar(texto, clave)
        entrada_resultado.config(state="normal")
        entrada_resultado.delete("1.0", tk.END)
        entrada_resultado.insert(tk.END, resultado)
        entrada_resultado.config(state="disabled")
    except ValueError:
        entrada_resultado.config(state="normal")
        entrada_resultado.delete("1.0", tk.END)
        entrada_resultado.insert(tk.END, "¡La clave debe ser un número entero!")
        entrada_resultado.config(state="disabled")

def ejecutar_descifrado():
    texto = entrada_texto.get("1.0", tk.END).strip()
    try:
        clave = int(entrada_clave.get())
        resultado = descifrar_cesar(texto, clave)
        entrada_resultado.config(state="normal")
        entrada_resultado.delete("1.0", tk.END)
        entrada_resultado.insert(tk.END, resultado)
        entrada_resultado.config(state="disabled")
    except ValueError:
        entrada_resultado.config(state="normal")
        entrada_resultado.delete("1.0", tk.END)
        entrada_resultado.insert(tk.END, "¡La clave debe ser un número entero!")
        entrada_resultado.config(state="disabled")


# --- 3. CONSTRUCCIÓN DE LA VENTANA GRÁFICA ---
ventana = tk.Tk()
ventana.title("Cifrador César")
ventana.geometry("480x420")
ventana.config(bg="#1e1e1e")  # Fondo oscuro elegante estilo VS Code

# Estilos reutilizables para las etiquetas
estilo_label = {"bg": "#1e1e1e", "fg": "#ffffff", "font": ("Calibri", 12, "bold")}

# Campo 1: Mensaje de entrada
tk.Label(ventana, text="Ingresa el mensaje", **estilo_label).pack(pady=(15, 5))
entrada_texto = tk.Text(ventana, height=4, width=50, bg="#2d2d2d", fg="#ffffff", insertbackground="white", font=("Consolas", 10))
entrada_texto.pack(pady=5)

# Campo 2: Clave numérica (k)
frame_clave = tk.Frame(ventana, bg="#1e1e1e")
frame_clave.pack(pady=10)
tk.Label(frame_clave, text="Clave (k): ", **estilo_label).pack(side=tk.LEFT)
entrada_clave = tk.Entry(frame_clave, width=8, bg="#2d2d2d", fg="#ffffff", insertbackground="white", font=("Consolas", 11), justify="center", show="•")
entrada_clave.pack(side=tk.LEFT)

# Campo 3: Resultado
tk.Label(ventana, text="Resultado", **estilo_label).pack(pady=(10, 5))
entrada_resultado = tk.Text(ventana, height=4, width=50, bg="#2d2d2d", fg="#00ffcc", font=("Consolas", 10), state="disabled")
entrada_resultado.pack(pady=5)

# Campo 4: Botones de acción (Cifrar / Descifrar)
# Fila 1: Botones de Cifrado y Descifrado
frame_botones_f1 = tk.Frame(ventana, bg="#1e1e1e")
frame_botones_f1.pack(pady=5)

btn_cifrar = tk.Button(frame_botones_f1, text="Cifrar", command=ejecutar_cifrado, width=12, bg="#007acc", fg="white", font=("Calibri", 11, "bold"), relief="flat")
btn_cifrar.pack(side=tk.LEFT, padx=8)

btn_descifrar = tk.Button(frame_botones_f1, text="Descifrar", command=ejecutar_descifrado, width=12, bg="#28a745", fg="white", font=("Calibri", 11, "bold"), relief="flat")
btn_descifrar.pack(side=tk.LEFT, padx=8)

# Fila 2: Botones de Limpiar y Copiar
frame_botones_f2 = tk.Frame(ventana, bg="#1e1e1e")
frame_botones_f2.pack(pady=8)

btn_limpiar = tk.Button(frame_botones_f2, text="Limpiar",command=limpiar_campos, width=12, bg="#d9534f", fg="white", font=("Calibri", 11, "bold"), relief="flat")
btn_limpiar.pack(side=tk.LEFT, padx=8)

btn_copiar = tk.Button(frame_botones_f2, text="Copiar", command=copiar_resultado, width=12, bg="#ffc107", fg="black", font=("Calibri", 11, "bold"), relief="flat")
btn_copiar.pack(side=tk.LEFT, padx=8)

# Fila 3: Cifrar Archivo Externo
frame_fila3 = tk.Frame(ventana, bg="#1e1e1e")
frame_fila3.pack(pady=6)

btn_archivo = tk.Button(frame_fila3, text="📁 Cifrar Archivo (Txt/Word/Pdf)", command=cifrar_archivo, width=27, bg="#6f42c1", fg="white", font=("Calibri", 11, "bold"), relief="flat")
btn_archivo.pack(side=tk.LEFT)

# Mantener la ventana abierta en ejecución
ventana.mainloop()